from __future__ import annotations

from copy import copy
from pathlib import Path

from openpyxl.styles import PatternFill
import pandas as pd


SCHEDULE_START = pd.Timestamp("2026-09-03 08:00")
SCHEDULE_END = pd.Timestamp("2026-09-04 08:00")


def demo_orders() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["UAT-260903-001", "SIM-C2-A", 600, "PCS", "急單", "2026-09-04", ""],
            ["UAT-260903-002", "SIM-C2-B", 480, "PCS", "一般", "2026-09-04", ""],
            ["UAT-260903-003", "SIM-MULTI-A", 360, "PCS", "一般", "2026-09-04", ""],
            ["UAT-260903-004", "SIM-C4-A", 720, "PCS", "急單", "2026-09-04", ""],
            ["UAT-260903-005", "SIM-C4-B", 540, "PCS", "一般", "2026-09-05", ""],
            ["UAT-260903-006", "SIM-MULTI-B", 420, "PCS", "低優先", "2026-09-05", ""],
            ["UAT-260903-007", "SIM-C5-A", 900, "PCS", "急單", "2026-09-04", ""],
            ["UAT-260903-008", "SIM-C5-B", 660, "PCS", "一般", "2026-09-05", ""],
            ["UAT-260903-009", "SIM-MULTI-C", 300, "PCS", "一般", "2026-09-04", "C2,C5"],
            ["UAT-260903-010", "SIM-C2-C", 840, "PCS", "低優先", "2026-09-05", ""],
        ],
        columns=["工單編號", "產品", "數量", "單位", "優先級", "交期", "工單限定機台"],
    )


def demo_rates() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["SIM-C2-A", "C2", 120, "C2-A"],
            ["SIM-C2-B", "C2", 96, "C2-B"],
            ["SIM-MULTI-A", "C2", 90, "MULTI-A"],
            ["SIM-MULTI-A", "C4", 84, "MULTI-A"],
            ["SIM-C4-A", "C4", 144, "C4-A"],
            ["SIM-C4-B", "C4", 108, "C4-B"],
            ["SIM-MULTI-B", "C4", 84, "MULTI-B"],
            ["SIM-MULTI-B", "C5", 78, "MULTI-B"],
            ["SIM-C5-A", "C5", 150, "C5-A"],
            ["SIM-C5-B", "C5", 110, "C5-B"],
            ["SIM-MULTI-C", "C5", 100, "MULTI-C"],
            ["SIM-MULTI-C", "C2", 92, "MULTI-C"],
            ["SIM-C2-C", "C2", 140, "C2-C"],
        ],
        columns=["產品", "機台", "產速_PCS_per_hr", "換模群組"],
    )


def demo_availability() -> pd.DataFrame:
    return pd.DataFrame(
        [["C2", SCHEDULE_START, SCHEDULE_END], ["C4", SCHEDULE_START, SCHEDULE_END], ["C5", SCHEDULE_START, SCHEDULE_END]],
        columns=["機台", "可用開始", "可用結束"],
    )


def demo_settings() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["排程開始", SCHEDULE_START],
            ["排程結束", SCHEDULE_END],
            ["時區", "Asia/Taipei"],
            ["允許跨班", "是"],
        ],
        columns=["設定項目", "設定值"],
    )


def demo_initial_state() -> pd.DataFrame:
    return pd.DataFrame(
        [["C2", "SIM-C2-A"], ["C4", "SIM-C4-A"], ["C5", "SIM-C5-A"]],
        columns=["機台", "初始產品"],
    )


def demo_changeovers() -> pd.DataFrame:
    return pd.DataFrame(
        [
            ["C2-A", "C2-B", 20],
            ["C2-B", "C2-A", 25],
            ["C2-A", "C2-C", 35],
            ["C2-C", "C2-A", 30],
            ["C4-A", "C4-B", 25],
            ["C4-B", "C4-A", 30],
            ["C5-A", "C5-B", 20],
            ["C5-B", "C5-A", 25],
            ["MULTI-A", "MULTI-B", 40],
            ["MULTI-B", "MULTI-C", 45],
            ["MULTI-C", "MULTI-A", 35],
        ],
        columns=["來源換模群組", "目標換模群組", "換模時間_分鐘"],
    )


def demo_workbook() -> dict[str, pd.DataFrame]:
    return {
        "待排工單": demo_orders(),
        "產品機台產速": demo_rates(),
        "機台初始狀態": demo_initial_state(),
        "換模時間": demo_changeovers(),
    }


def write_demo_excel(path: str | Path) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for sheet_name, frame in demo_workbook().items():
            frame.to_excel(writer, sheet_name=sheet_name, index=False)
        required_columns = {
            "待排工單": {"工單編號", "產品", "數量", "單位", "優先級", "交期"},
            "產品機台產速": {"產品", "機台", "產速_PCS_per_hr", "換模群組"},
            "機台初始狀態": {"機台", "初始產品"},
            "換模時間": {"來源換模群組", "目標換模群組", "換模時間_分鐘"},
        }
        required_fill = PatternFill("solid", fgColor="FCE4EC")
        header_fill = PatternFill("solid", fgColor="E5E7EB")
        for worksheet in writer.book.worksheets:
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                font = copy(cell.font)
                font.bold = True
                cell.font = font
                cell.fill = required_fill if cell.value in required_columns.get(worksheet.title, set()) else header_fill
            for column_cells in worksheet.columns:
                values = [str(cell.value) if cell.value is not None else "" for cell in column_cells]
                width = min(max(max(len(value) for value in values) + 2, 12), 24)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
    return path
